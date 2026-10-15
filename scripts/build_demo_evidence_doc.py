"""使用Codex捆绑Python与python-docx生成本轮模拟案例及实验报告。"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '航旅智策应用/runtime/evidence'


def main():
    r = json.loads((OUT / 'results.json').read_text(encoding='utf-8'))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.1), Cm(1.9)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(.85)
    sec.different_first_page_header_footer = False
    doc.settings.odd_and_even_pages_header_footer = False
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        style = doc.styles[name]
        style.font.name = '宋体' if name == 'Normal' else '黑体'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), style.font.name)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
    doc.styles['Heading 2'].font.size = Pt(12)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].paragraph_format.space_after = Pt(5)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = '航旅智策模拟案例与对照实验报告'
    header = sec.header.paragraphs[0]
    header.text = '华翼朝阳  航旅智策项目演示证据'
    header.style = 'Caption'
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('模拟数据验证  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(9)

    def p(text, style=None):
        return doc.add_paragraph(text, style)

    def head(text):
        doc.add_heading(text, level=1)

    def page(title):
        doc.add_page_break()
        head(title)

    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        widths = widths or [16.6 / len(headers)] * len(headers)
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = str(text)
        repeat = OxmlElement('w:tblHeader')
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            border = OxmlElement('w:' + side)
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '4')
            border.set(qn('w:color'), 'D9D9D9')
            borders.append(border)
        t._tbl.tblPr.append(borders)
        for ri, row in enumerate(t.rows):
            no_split = OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(widths[ci])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8ECF5' if ri == 0 else 'F6F8FC' if ri % 2 == 0 else 'FFFFFF')
                props.append(shade)
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'bottom', 'left', 'right']:
                    item = OxmlElement('w:' + side)
                    item.set(qn('w:w'), '85')
                    item.set(qn('w:type'), 'dxa')
                    margins.append(item)
                props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.15
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.bold = ri == 0
                        run.font.name = '宋体'
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        p('')
        return t

    def picture(name, caption, width=16):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(OUT / 'figures' / name), width=Cm(width))
        p(caption, 'Caption')

    def ci(scope, baseline='近30天频次', metric='Recall@3'):
        return next(x for x in r['ranking']['results'][scope]['paired_differences'] if x['baseline'] == baseline and x['metric'] == metric)

    def pp(x):
        return f'{x * 100:.3f}'

    c = r['cases']
    national = r['ranking']['results']['国庆专项']
    frozen = r['ranking']['results']['固定测试']
    pos, neg = c['positive'], c['deferred']
    d = r['decision']['results']['未退订全量同群']
    delta = ci('国庆专项')

    doc.add_paragraph('航旅智策模拟案例与对照实验报告', 'Title')
    p('完整案例 动态变化与公平对照', 'Subtitle')
    p('执行日期 ' + r['created_at'][:10] + '    源码版本 ' + r['source_commit'][:7])
    p('本报告用于项目报告中的原型验证和创新实践分析。已实际完成两名模拟旅客的营销流程、同一旅客的三个时间截点回放，以及需求排序与业务决策两组对照。结果支持当前模拟环境中的流程可行性与局部方法改善，不支持真实购票转化、客座率、收入或ROI提升。')
    head('主要结果')
    table(['验证内容', '实际结果', '可支持的结论'], [
        ['完整流程', '1人模拟触达 1人暂缓\n重复执行新增0条', '审核和资格门控可运行'],
        ['动态变化', 'SIMU_00748\n长沙转为成都', '画像可随已观察证据更新'],
        ['国庆排序', f'Recall@3 {national["metrics"][0]["Recall@3"]:.4f}\n较频次规则提高{pp(delta["difference"])}个百分点', '本模拟样本上排序改善'],
        ['业务条件', '全量未退订同群127人增至150人\n共同前300人仍为4人', '收益依赖名单分布']], [3.0, 7.0, 6.6])
    head('证据边界')
    p('行为与营销反馈均为模拟数据，航班为2024历史样本；不具备实时可售库存。文案为本地模板，海报为SVG合成，本次没有部署或调用本地大模型。审核确认由实验脚本提交，验证的是审核门控，不是独立人工内容质量评价。')
    p('原始三份Excel、训练产物及用户活动数据库在执行前后通过指纹核验保持不变。实验活动写入独立状态库。问卷证据没有参与本轮实验，不将其与万人模拟旅客混合。')

    page('实验设计与公平性')
    table(['阶段', '观察截点', '用途'], [
        ['训练', '2024年1至6月 每月15日', '沿用已训练模型 不重新拟合'],
        ['验证', '2024年7至9月 每月15日', '模型选择标签截至9月22日'],
        ['固定测试', '10月15日 11月15日 12月15日', '30000个旅客快照 其中29998有效'],
        ['国庆专项', '9月30日', '预测10月1至7日 9998有效快照']], [2.5, 7.1, 7.0])
    p('同一批10000名既有旅客、同一41个目的地、同一观察截点和未来7天标签用于所有排序方法。无相关目的地的快照单列，不计入Recall与NDCG均值；并列分数按同一目的地字典稳定排序。既有旅客跨阶段出现，结果不等同于新用户泛化。')
    table(['方法', '可用信息与规则'], [
        ['动态学习模型', '已冻结梯度提升 23项截点前特征'],
        ['逻辑回归', '既有对照模型 使用同一截点特征'],
        ['近30天频次', '近30天目的地搜索过程占比'],
        ['14天时间衰减', '仅对已观察的过程按时间衰减'],
        ['最近关注', '按最近关注距离排序'],
        ['固定六月画像', '6月30日画像固定后不更新'],
        ['训练热门', '只使用训练期的目的地平均频次']], [4.1, 12.5])
    p('Recall@3衡量前三个推荐覆盖了多少未来相关目的地；NDCG@3同时考虑相关方向的排序位置。相关定义为后7天同目的地至少两次搜索过程，不是购买。模型与频次规则为预指定主要对照；其他比较为探索性。')
    p('配对置信区间按旅客整组重采样2000次，固定种子20261015；同一旅客跨日期的快照保持成组，并按有效快照数加权。区间仅反映当前模拟样本中的抽样变化，不覆盖数据生成机制的不确定性。')

    page('完整营销案例与暂缓案例')
    table(['旅客', '方向', '日期与预算', '处理'], [
        [pos['user_id'], pos['usual_departure_city'] + ' → ' + pos['destination'], '10月1日明确计划\n机票上限1550元', '历史CA3989 样本价880元\n本地模拟触达1条'],
        [neg['user_id'], neg['usual_departure_city'] + ' → ' + neg['destination'], '日期未确认\n本次预算未知', '素材可制作\n业务条件不足而暂缓']], [3.0, 3.8, 4.6, 5.2])
    p('案例选择只依赖9月30日前的证据和执行条件，不根据未来结果挑选。成功案例选择用于说明完整流程，暂缓案例用于说明高关注分数不能替代明确出行条件。')
    for text in [
        '冻结名单 以国庆场景从当前候选名单保留两名指定模拟旅客，活动名单版本为2。',
        '生成素材 两条模板文案和两份SVG海报完成事实槽位合成；未使用大模型。',
        '审核门控 脚本核对风险标记后提交事实、历史价格、模拟属性三项确认；来源记录与操作审计保留。',
        '执行检查 日期、预算、历史候选、退订、时段、频控与素材均重新检查，仅SIMU_00179通过。',
        '重复执行 第二次请求新增0条，历史唯一触达记录保留，另一旅客仍暂缓。',
        '后续复盘 只读取10月1至7日的原模拟行为，不回填此前画像，也不将后续行为归因于新触达。']:
        p(text)
    outcomes = c['evaluation']['outcomes']
    table(['观察对象', '持续关注', '机票查询', '收藏动作', '任意关注'], [[
        '已模拟触达的1人', outcomes.get('continue', 0), outcomes.get('fare', 0), outcomes.get('favorite', 0), outcomes.get('attention', 0)]], [4.6, 3, 3, 3, 3])
    p('新活动记录中的exposure、click、payment均为null。上述0或1仅表示既有模拟行为中是否满足代理标签；不是新活动点击率、支付率或因果效果。')
    p('素材、完整活动JSON、执行账本和操作截图随材料交付。重复执行后活动状态可能显示“暂不触达”，它表示该次请求没有新增记录，不抹除首次触达。')

    page('同一旅客的动态变化')
    rows = []
    for s in c['dynamic']['snapshots']:
        shares = {x['name']: x['share'] for x in s['profile']['destinations']}
        rows.append([s['date'], f'{shares["长沙"]:.1f}%', f'{shares["成都"]:.1f}%', s['model_top1']])
    table(['观察截点', '长沙关注占比', '成都关注占比', '模型首选'], rows, [4.2, 4.1, 4.1, 4.2])
    picture('dynamic-profile.png', '图1  SIMU_00748的时间衰减关注份额。份额不等于购票概率。')
    p('该旅客出发地为太原。9月25日的成都自由行攻略、9月26日的成都酒店怎么选，以及此前记录共同影响近30天时间衰减画像。窗口滚动和旧行为衰减也会改变份额，不能把变化全部归于某一次新增搜索。')
    p('9月30日最后一条可见搜索仍为长沙旅游，来源SIMB_00748_304，但模型首选仍为成都。这说明当前机制综合近期证据，并非直接把最后一次搜索当作整体出游意向。三个截点均晚于模型选择标签结束日期9月22日。')
    p('动态示例按已观察的画像与模型首选变化筛选，作用是解释机制。它不能代替全量对照，也不证明该旅客后来真实前往成都。交互演示页可拖动截点，逐条查看来源记录。')

    page('需求排序对照的实际结果')
    table(['方法', '固定测试\nRecall@3', '固定测试\nNDCG@3', '国庆专项\nRecall@3', '国庆专项\nNDCG@3'], [
        [a['method'], f'{a["Recall@3"]:.4f}', f'{a["NDCG@3"]:.4f}', f'{b["Recall@3"]:.4f}', f'{b["NDCG@3"]:.4f}']
        for a, b in zip(frozen['metrics'], national['metrics'])], [4.6, 3, 3, 3, 3])
    table(['主要对照', 'Recall@3差值\n百分点', '配对95%区间\n百分点'], [
        [scope + ' 对频次规则', pp(ci(scope)['difference']), ' 至 '.join(pp(v) for v in ci(scope)['ci95'])]
        for scope in ['固定测试', '国庆专项']], [7.0, 4.2, 5.4])
    picture('ranking-differences.png', '图2  动态学习模型相对强规则基线的差值。其他比较为探索性，未进行多重比较调整。', 15.6)
    p('固定测试的七种方法指标与已有训练报告逐项复现一致，绝对误差小于1e-7。国庆专项单独执行，未为改善结果而重训模型或修改阈值。稳定模拟习惯可能使模型和规则均表现较高，不宜向真实航空用户外推。')

    page('业务约束的收益与零改善解释')
    table(['范围与方法', '固定分母', '历史匹配', '业务条件满足', '暂缓'], [
        [scope + '\n' + row['method'], row['users'], row['historical_matches'], row['business_ready'], row['deferred']]
        for scope, result in r['decision']['results'].items() for row in result['metrics']], [6.2, 2.6, 2.6, 2.6, 2.6])
    for key, label in [('matched', '历史匹配率'), ('ready', '业务条件满足率')]:
        value = d['paired_differences'][key]
        p(f'全量未退订同群的{label}差值为{pp(value["difference"])}个百分点，配对95%区间为' + '至'.join(pp(v) for v in value['ci95']) + '个百分点。')
    diag = r['decision']['results']['共同兴趣名单300人']['diagnostics']
    p('共同前300人为什么没有改善：' + '；'.join(f'{k}{v}人' for k, v in diag.items()) + '。这300人中没有发生备选目的地替换，因而匹配人数和业务条件满足人数保持不变。约束不能创造缺失的历史航班或明确预算。')
    picture('decision-comparison.png', '图3  全量同群有局部改善，但共同前300人名单没有改善。')
    p('该消融固定同一模型、同一近期兴趣候选和预算日期规则，只改变候选选择方式。共同300人按首选分数冻结，不前筛航班。产品当前默认先筛历史航班再截取的300人名单与其不同，不能用产品名单300/300覆盖率证明本对照获益。业务对照仅离线，不执行不合格建议。')

    page('名单选择策略的探索性优化')
    policy = r['decision']['selection_policy_exploration']
    p('前300人没有候选替换收益，不意味着整个名单选择流程不能优化。另一问题是在同一可达旅客池、同一名额和费用上限下，把有限名额分给哪些人。该试验允许三种方法选中不同旅客，与上一页固定同一名单的候选消融分开报告。')
    table(['名单策略', '实际选中', '配额缺口', '历史匹配', '业务合格', '后续持续关注'], [
        [x['policy'], x['selected'], x['shortfall'], x['historical_matches'], x['business_ready'], x['future_continue']]
        for x in policy['metrics']], [5.1, 2.3, 2.3, 2.3, 2.3, 2.3])
    p('三组均从9799名未退订旅客中选择，最多300人，单次情景费用0.5元，总费用上限150元。兴趣分数策略按首选分数选人；历史候选优先策略先筛选有历史匹配的旅客，再按选择后分数排序；业务资格优先策略先筛选当次计划、预算和产品条件均满足的旅客，再排序。所有选择只使用9月30日前可见信息。')
    p('推荐的应用优化是将“产品条件优先”和“业务资格优先”作为可切换的名单目标，显示实际选中人数、缺额及暂缓原因。业务资格优先本轮只能选出150人，不能为了凑满300人而提高预算、改写日期或补造航班。其他名额可用于邀请旅客补充条件，但这类需求确认任务应与航线营销触达分别计量。')
    table(['名单策略', '名单内业务合格率', '每300个名额的合格率', '合格任务情景费用'], [
        [x['policy'], f'{x["ready_rate_in_selected"]*100:.2f}%', f'{x["ready_per_300_capacity"]*100:.2f}%', f'{x["qualified_cost_scenario"]:.2f}元']
        for x in policy['metrics']], [6.1, 3.5, 3.5, 3.5])
    p('后续持续关注在名单冻结后才评价，不用于选人。不同策略的样本构成不同，应同时报告人数、比率、缺额和后续代理表现；不能只用资格优先名单的100%合格率掩盖名单不足额。实际送达和转化仍未验证。')
    rates = [f'{x["future_continue"]}/{x["selected"]} 即{x["future_continue_rate"]*100:.2f}%' for x in policy['metrics']]
    p('三种策略的后续持续关注率依次为' + '、'.join(rates) + '。资格优先在业务条件上更完整，但本次关注代理率反而低于原兴趣名单，说明业务可行性与高关注不是同一个目标，不应只挑选有利指标。')
    p('本试验根据零改善诊断提出，属于探索性业务策略分析，不是事先预注册的确认试验。当前网站默认已采用历史候选优先；业务资格优先在此为离线实验，本轮没有改动产品默认逻辑。进一步部署前宜在新的时间窗口验证其稳定性。')

    page('复现方法与证据文件')
    p('在仓库根目录执行以下命令。采用已核验的原始工作簿和训练产物，不下载或启动语言模型。统计图另需matplotlib，演示视频另需imageio-ffmpeg；依赖下载继承已配置HTTPS代理。')
    p('航旅智策应用\\.venv\\Scripts\\python.exe scripts\\run_demo_evidence.py')
    p('航旅智策应用\\.venv\\Scripts\\python.exe scripts\\build_demo_evidence_view.py')
    p('航旅智策应用\\.venv\\Scripts\\python.exe -m unittest scripts.test_demo_evidence -v')
    table(['文件', '用途'], [
        ['results.json', '协议 主要结果 区间 案例 指纹与隔离核验'],
        ['ranking-frozen.csv ranking-national.csv', '两范围七种方法的完整结果'],
        ['ranking-*-by-date.csv', '逐日期明细 防止只展示有利截点'],
        ['per-user-frozen.npz per-user-national.npz', '配对重采样所用逐人指标数组'],
        ['decision-paired-users.csv', '同一旅客两策略结果与候选诊断'],
        ['selection-policy-comparison.csv', '同预算名单策略结果 另附名单成员CSV'],
        ['workflow-campaign.json state目录', '独立活动与审核记录 仅模拟案例状态'],
        ['案例与实验演示.html figures目录', '离线交互演示及可分享PNG和SVG统计图']], [6.8, 9.8])
    p('三项统计检查覆盖稳定并列、无相关标签的单列处理、跨日期旅客整组加权及配对有效性。运行脚本还核对固定测试复现、首次触达唯一记录、重复触达零新增、未知点击支付字段，以及输入和用户状态未修改。')
    p('数据版本 ' + r['source_version'] + '。随机种子20261015；2000次配对重采样。环境 Python ' + r['environment']['python'] + '，NumPy ' + r['environment']['numpy'] + '，scikit-learn ' + r['environment']['scikit_learn'] + '。')
    for name, sha in r['source_sha256'].items():
        p(name + '  SHA256前16位 ' + sha[:16])
    p('完整SHA256、当前关键源码指纹及模型产物指纹见results.json。原始行为生成程序及生成参数尚未在本轮独立审计；样本数大不能替代真实数据验证。')

    page('创新分析与演示使用')
    p('可支持的创新论点是把动态需求识别与当次业务约束连接成可追溯的营销决策流程。排序实验说明当前模拟样本上的增益，案例说明行为证据如何转成航线与内容，业务消融说明部分备选需求可被恢复，同时保留名单无改善的边界。使用通用梯度提升和聚类不等同于原创算法突破。')
    p('建议以营销人员端为演示主线，个人旅行定制作为偏好收集与服务承接入口。原型已能服务方案说明；真实转化、客户满意度和收入需要企业授权数据与后续试点。渠道及时间预测只刻画访问习惯，不证明最优触达策略。')
    table(['讲解时间建议', '画面与讲解重点'], [
        ['0至15秒', '概览说明数据属性与三个验证目标'],
        ['15至40秒', '成功与暂缓旅客的日期 预算 历史候选'],
        ['40至65秒', '模板素材 审核门控 首次与重复触达'],
        ['65至90秒', '三个观察截点 来源记录与首选变化'],
        ['90至115秒', '排序对照和配对区间 不只比较热门基线'],
        ['115至140秒', '全量改善与共同300人零改善的区别'],
        ['140至150秒', '边界及真实小规模试点计划']], [3.6, 13.0])
    p('下一步优先采集自愿参与者的前后两次偏好与任务操作记录，验证未知意向、预算变化和动态更新；独立记录人工完成任务的耗时与修改次数。本轮不把这类待开展工作写成已完成实验。若最终参赛强调AIGC，应另补真实生成样本和事实一致性评价，本次仍为模板演示。')
    p('资料来源为项目原始模拟工作簿、冻结训练报告、当前可运行源码与本轮实际实验结果。实验脚本和文稿由AI辅助整理，数值来自本机执行。操作截图及视频只展示模拟活动，不包含真实身份或实际外部营销。')
    # The bundled seed may contain title borders or first/even-page furniture.
    for element in [doc._element, doc.styles.element]:
        for border in list(element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for section in doc.sections:
        for part in [section.header, section.first_page_header, section.even_page_header]:
            for paragraph in part.paragraphs:
                paragraph.text = '华翼朝阳  航旅智策项目演示证据'
                paragraph.style = 'Caption'
        for part in [section.footer, section.first_page_footer, section.even_page_footer]:
            paragraph=part.paragraphs[0]
            paragraph.clear()
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run('模拟数据验证  |  ')
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');paragraph._p.append(field)
            for run in paragraph.runs:run.font.size=Pt(9)
    doc.save(OUT / '航旅智策模拟案例与对照实验报告.docx')
    print('WORD_REPORT_READY', OUT / '航旅智策模拟案例与对照实验报告.docx')


if __name__ == '__main__':
    main()
