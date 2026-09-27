"""将可编辑实施指南生成Word。需要python-docx；不属于Demo运行依赖。"""
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'project-plan.md'
TARGET = ROOT / 'deliverables' / '航旅数字画像项目实施指南.docx'


def font(run, size=11, bold=False, family='宋体'):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), family)


def text(paragraph, content, size=11):
    for index, part in enumerate(re.split(r'\*\*(.*?)\*\*', content)):
        font(paragraph.add_run(part), size, bool(index % 2))


def table(doc, lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', cell) for cell in row)]
    result = doc.add_table(rows=0, cols=len(rows[0]))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    proportions = {2: [.40,.60], 3: [.24,.40,.36], 4: [.22,.25,.26,.27]}.get(len(rows[0]))
    if rows[0][0] == '编号': proportions = [.08,.50,.42]
    if rows[0][0] == '文件': proportions = [.56,.44]
    for col, fraction in zip(result.columns, proportions):
        col.width = Cm(16.6 * fraction)
    for index, data in enumerate(rows):
        cells = result.add_row().cells
        for cell, value, fraction in zip(cells, data, proportions):
            cell.width = Cm(16.6 * fraction)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'DAE8F3' if index == 0 else ('F5F8FA' if index % 2 == 0 else 'FFFFFF'))
            properties.append(shading)
            margins = OxmlElement('w:tcMar')
            for name, val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                node = OxmlElement('w:' + name); node.set(qn('w:w'),val); node.set(qn('w:type'),'dxa'); margins.append(node)
            properties.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.2
            font(paragraph.add_run(value), 9.5, index == 0)
        no_split = OxmlElement('w:cantSplit'); result.rows[-1]._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); result.rows[-1]._tr.get_or_add_trPr().append(repeat)
    borders = OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        edge = OxmlElement('w:' + name)
        for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]: edge.set(qn('w:' + key), val)
        borders.append(edge)
    result._tbl.tblPr.append(borders)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    font(spacer.add_run(''), 3)


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
    section.top_margin = Cm(1.9); section.bottom_margin = Cm(1.9)
    normal = doc.styles['Normal']
    normal.font.name = '宋体'; normal.font.size = Pt(11)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.line_spacing = 1.30
    normal.paragraph_format.space_after = Pt(5)
    for name,size in [('Title',24),('Heading 1',17),('Heading 2',12.5)]:
        style = doc.styles[name]
        style.font.name = '黑体'; style.font.size = Pt(size); style.font.bold = name != 'Title'
        style.font.color.rgb = RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
        style.paragraph_format.space_before = Pt(12 if name == 'Heading 2' else 0)
        style.paragraph_format.space_after = Pt(10)
    doc.core_properties.title = '航旅数字画像项目实施指南'
    doc.core_properties.subject = '钻石命题2的研究框架与实施步骤'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    # 清除Word内置Title等样式可能继承的装饰线。
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(foot.add_run('第 '),9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
    foot._p.append(field)
    font(foot.add_run(' 页'),9)
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0; first_title = True; new_page = False
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line == '<!-- page -->':
            new_page = True; i += 1; continue
        if line.startswith('|'):
            group = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                group.append(lines[i]); i += 1
            table(doc,group); continue
        if line.startswith('# '):
            heading = doc.add_paragraph(line[2:],style='Title' if first_title else 'Heading 1')
            heading.paragraph_format.page_break_before = new_page
            new_page = False
            first_title = False
        elif line.startswith('## '):
            doc.add_paragraph(line[3:],style='Heading 2')
        else:
            p = doc.add_paragraph()
            text(p,line)
        i += 1
    TARGET.parent.mkdir(parents=True,exist_ok=True)
    doc.save(TARGET)
    print(TARGET)


if __name__ == '__main__':
    main()
